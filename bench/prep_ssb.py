"""One-time offline prep for tasks_ssb.py: SpreadsheetBench Verified (400) workbooks to plain JSON.

Not used at run time; needs openpyxl (`uv run --with openpyxl python prep_ssb.py`). tasks_ssb.py reads only the
JSON this writes, with the standard library.

Source: https://huggingface.co/datasets/KAKA22/SpreadsheetBench, file spreadsheetbench_verified_400.tar.gz
(CC BY-SA 4.0, see bench/ext/SSB_LICENSE). Downloaded to bench/ext/ if missing and unpacked to
bench/ext/ssb_raw/. The archive holds dataset.json (id, instruction, instruction_type, answer_position,
spreadsheet_path; some entries also answer_sheet, data_position, exclude) and one folder per id with
<n>_<id>_init.xlsx (input) and <n>_<id>_golden.xlsx (answer) per test case n. Every Verified folder holds
exactly one pair (n = 1); five use the names initial.xlsx / golden.xlsx, and 42930's golden file is named
1_43930_golden.xlsx (see test_case_files).

Output: bench/ext/ssb/<id>.json with
  id, instruction, instruction_type, answer_position, exclude (if the dataset marks it),
  sheetnames            sheets of the golden workbook (official rule: a range without a sheet is on the first one)
  answer                [[sheet, "A1:B3"], ...] parsed from answer_position, or position_error
  test_cases            per test case n:
      sheets            [{name, hidden, max_row, max_col, rows}] of the INPUT workbook; rows[r][c] is the value of
                        the cell at row r+1, column c+1, over A1 to the last non-empty cell. A sheet over
                        MAX_CELLS cells keeps only its size (too_large) and no rows.
      formula_cells     number of formula cells in the input sheets
      uncached          input formula cells left without a value (see below), as "Sheet!A1"
      input_from_golden number of input formula cells whose value came from the golden file (see below)
      expected          {"Sheet!A1": value} for every cell of the answer position, from the golden workbook
      expected_uncached golden answer cells that hold a formula without a cached value
      init_answer       {"Sheet!A1": value} of the same cells in the input workbook
      answer_sheet_missing  answer sheets absent from the input workbook
      answer_too_large  set (and expected left out) when the answer position spans over MAX_CELLS cells
      volatile          an input formula or golden answer formula calls TODAY, NOW or RAND*

Values: str, int, float, bool or null as in openpyxl with data_only=True (the cached value Excel saved for a
formula cell); datetime as {"date": "YYYY-MM-DDTHH:MM:SS"}, time as {"time": "HH:MM:SS"}, timedelta as
{"duration_s": seconds}.

Formulas: no formula is evaluated here. A formula cell takes the value Excel cached in the file (openpyxl
data_only=True). An input formula cell without a cached value takes the golden file's cached value for the same
cell when the golden cell holds the same formula text; otherwise it is listed in uncached. A golden answer cell
whose formula has no cached value is listed in expected_uncached. tasks_ssb.py leaves out instances with either.
Recalculating with LibreOffice (what the official open_spreadsheet.py does before evaluation) was tried and not
used. On the 68 instances with an uncached formula in the input or golden file, LibreOffice's recalculation of
the golden file differed from Excel's cached value on 1,357 of the 7,935 formula cells Excel had cached: 1,239 in
one file (#REF! read as #NAME?), the other 118 in six files with array formulas (other values, or none). Its
values for the uncached cells therefore cannot be trusted as gold.
"""
import datetime
import json
import re
import sys
import tarfile
import urllib.request
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

EXT = Path(__file__).resolve().parent / "ext"
URL = "https://huggingface.co/datasets/KAKA22/SpreadsheetBench/resolve/main/spreadsheetbench_verified_400.tar.gz"
TAR = EXT / "spreadsheetbench_verified_400.tar.gz"
RAW = EXT / "ssb_raw"
OUT = EXT / "ssb"
MAX_CELLS = 40000
_RANGE = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?$")


def enc(v):
    if isinstance(v, datetime.datetime):
        return {"date": v.isoformat(timespec="seconds")}
    if isinstance(v, datetime.date):
        return {"date": datetime.datetime(v.year, v.month, v.day).isoformat(timespec="seconds")}
    if isinstance(v, datetime.time):
        return {"time": v.replace(microsecond=0).isoformat()}
    if isinstance(v, datetime.timedelta):
        return {"duration_s": v.total_seconds()}
    if v is None or isinstance(v, (str, bool, int, float)):
        return v
    return str(v)


def is_formula(cell):
    return cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))


def formula_text(cell):
    v = cell.value
    return getattr(v, "text", v)              # ArrayFormula keeps its text in .text


def parse_position(pos, sheetnames):
    """answer_position to [(sheet, "A1:B3")]. The official evaluation splits on "," and "!" and strips quotes;
    this does the same but also repairs the malformed entries in the data (a quote on one side only, the quote
    after the "!", a sheet name containing a comma). Raises ValueError for anything else (whole columns, a
    sheet that does not exist, a range without a column)."""
    pieces, buf = [], ""
    for p in pos.split(","):
        buf = f"{buf},{p}" if buf else p
        if "!" in buf or _RANGE.match(buf.strip().strip("'").strip()):
            pieces.append(buf)
            buf = ""
    if buf:
        raise ValueError(f"cannot parse {buf!r}")
    lower = {s.lower(): s for s in sheetnames}
    out = []
    for p in pieces:
        if "!" in p:
            sheet, rng = p.rsplit("!", 1)
            sheet = sheet.strip().strip("'").strip()
        else:
            sheet, rng = sheetnames[0], p
        rng = rng.strip().strip("'").strip().replace("$", "")
        if sheet.lower() not in lower:
            raise ValueError(f"no sheet {sheet!r}")
        if not _RANGE.match(rng):
            raise ValueError(f"bad range {rng!r}")
        out.append((lower[sheet.lower()], rng))
    return out


def cells_of(rng):
    m = _RANGE.match(rng)
    c1, r1 = column_index_from_string(m[1]), int(m[2])
    c2, r2 = (column_index_from_string(m[3]), int(m[4])) if m[3] else (c1, r1)
    for r in range(min(r1, r2), max(r1, r2) + 1):
        for c in range(min(c1, c2), max(c1, c2) + 1):
            yield f"{get_column_letter(c)}{r}"


def used_bounds(ws_f):
    """Last row and column that hold a value or formula (ws.max_row also counts formatted empty cells)."""
    mr = mc = 0
    for row in ws_f.iter_rows():
        for c in row:
            if c.value is not None and c.value != "":
                mr, mc = max(mr, c.row), max(mc, c.column)
    return mr, mc


_VOLATILE = re.compile(r"\b(?:TODAY|NOW|RAND|RANDBETWEEN|RANDARRAY)\s*\(", re.IGNORECASE)


def _ftext(cell):
    return str(formula_text(cell)) if is_formula(cell) else ""


def test_case(init, golden, answer):
    wf = openpyxl.load_workbook(init, data_only=False)
    wv = openpyxl.load_workbook(init, data_only=True)
    gf = openpyxl.load_workbook(golden, data_only=False)
    gv = openpyxl.load_workbook(golden, data_only=True)
    sheets, uncached, from_golden, nform, volatile = [], [], 0, 0, False
    for ws_f in wf.worksheets:
        ws_v = wv[ws_f.title]
        mr, mc = used_bounds(ws_f)
        entry = {"name": ws_f.title, "hidden": ws_f.sheet_state != "visible", "max_row": mr, "max_col": mc}
        if mr * mc > MAX_CELLS:
            entry["too_large"] = True
            sheets.append(entry)
            continue
        rows = []
        for r in range(1, mr + 1):
            row = []
            for c in range(1, mc + 1):
                cf, v = ws_f.cell(r, c), ws_v.cell(r, c).value
                if is_formula(cf):
                    nform += 1
                    volatile = volatile or bool(_VOLATILE.search(_ftext(cf)))
                    key = f"{ws_f.title}!{cf.coordinate}"
                    if v is None:
                        g = gf[ws_f.title].cell(r, c) if ws_f.title in gf.sheetnames else None
                        gval = gv[ws_f.title].cell(r, c).value if g is not None else None
                        if g is not None and _ftext(g) == _ftext(cf) and gval is not None:
                            v, from_golden = gval, from_golden + 1
                        else:
                            uncached.append(key)
                row.append(enc(v))
            rows.append(row)
        entry["rows"] = rows
        sheets.append(entry)
    out = {"sheets": sheets, "formula_cells": nform, "uncached": uncached, "input_from_golden": from_golden}
    if sum(1 for _, rng in answer for _ in cells_of(rng)) > MAX_CELLS:
        return {**out, "answer_too_large": True}
    expected, exp_unc, init_ans, missing = {}, [], {}, []
    for sheet, rng in answer:
        if sheet not in wf.sheetnames:
            missing.append(sheet)
        for a in cells_of(rng):
            key = f"{sheet}!{a}"
            gcell = gf[sheet][a]
            gval = gv[sheet][a].value
            volatile = volatile or bool(_VOLATILE.search(_ftext(gcell)))
            if is_formula(gcell) and gval is None:
                exp_unc.append(key)
            expected[key] = enc(gval)
            init_ans[key] = enc(wv[sheet][a].value) if sheet in wv.sheetnames else None
    return {**out, "expected": expected, "expected_uncached": exp_unc,
            "init_answer": init_ans, "answer_sheet_missing": sorted(set(missing)), "volatile": volatile}


def test_case_files(folder):
    """[(n, input, golden)]. Most folders use <n>_<id>_init.xlsx / <n>_<id>_golden.xlsx; a few use
    initial.xlsx / golden.xlsx, and one golden file carries a mistyped id. A folder with one input and one
    golden file is one test case whatever the names; otherwise files are paired by their leading number."""
    inits = sorted(p for p in folder.glob("*.xlsx") if "init" in p.name)
    golds = sorted(p for p in folder.glob("*.xlsx") if "golden" in p.name)
    if len(inits) == 1 and len(golds) == 1:
        return [(1, inits[0], golds[0])]
    by_no = {int(p.name.split("_")[0]): p for p in golds}
    return [(int(p.name.split("_")[0]), p, by_no[int(p.name.split("_")[0])]) for p in inits]


def main():
    EXT.mkdir(exist_ok=True)
    if not TAR.exists():
        print("downloading", URL)
        urllib.request.urlretrieve(URL, TAR)
    root = RAW / "spreadsheetbench_verified_400"
    if not root.exists():
        RAW.mkdir(exist_ok=True)
        with tarfile.open(TAR) as t:
            t.extractall(RAW, filter="data")
    OUT.mkdir(exist_ok=True)
    data = json.loads((root / "dataset.json").read_text(encoding="utf-8"))
    only = set(sys.argv[1:])
    for i, d in enumerate(data):
        sid = str(d["id"])
        if only and sid not in only:
            continue
        folder = root / d["spreadsheet_path"]
        rec = {k: d[k] for k in ("instruction", "instruction_type", "answer_position") if k in d}
        rec = {"id": sid, **rec}
        if "exclude" in d:
            rec["exclude"] = d["exclude"]
        try:
            cases = test_case_files(folder)
            rec["sheetnames"] = openpyxl.load_workbook(cases[0][2], read_only=True).sheetnames
            try:
                rec["answer"] = parse_position(d["answer_position"], rec["sheetnames"])
            except ValueError as e:
                rec["position_error"] = str(e)
            if "answer" in rec:
                rec["test_cases"] = [{"no": n, **test_case(inp, g, rec["answer"])} for n, inp, g in cases]
        except Exception as e:                      # unreadable workbook
            rec["load_error"] = f"{type(e).__name__}: {e}"
        (OUT / f"{sid}.json").write_text(json.dumps(rec, ensure_ascii=False), encoding="utf-8")
        print(f"{i + 1}/{len(data)} {sid}", rec.get("position_error") or rec.get("load_error") or "", flush=True)


if __name__ == "__main__":
    main()
